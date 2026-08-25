#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
L.U.N.A 정식 빌드 — 저작 xlsx 2파일 → 검증 → json/(도메인별 분할 배포)

  사용:  python3 데이터/tools/build.py                # 기본: LUNA_System.xlsx + LUNA_Narrative.xlsx
         python3 데이터/tools/build.py <xlsx…>        # 다른 파일(들) 지정
         python3 데이터/tools/build.py --strict       # 영어 누락 = 에러 (번역 완료 후 사용)

  v2.0 확정: 저작은 '수정 권한' 축으로 2파일 분리 (xlsx는 git 머지 불가 → 파일 1개 = 소유자 1명)
    LUNA_System.xlsx    — 제조·운영: 칵테일/선반/성격/취향/밸런스/RandomWaves
    LUNA_Narrative.xlsx — 대사·서사: 대본/Barks/단골/퀘스트/연출/캐스트/RegularSlots
  규칙: 한 시트는 정확히 한 파일에만 존재한다. 두 파일에 같은 시트가 있으면 빌드 에러
        (참고용 복사가 병합돼 유령 데이터가 되는 사고 차단).
  배포는 구엔진처럼 도메인별 JSON 분할. 검증 실패 시 JSON을 내보내지 않는다.
  gen_luna_data.py는 "코드 시드 → 시트 재생성" 전용. 팀 저작 시작 후엔 build.py만 사용(gen은 시트를 덮는다).
"""
import hashlib
import json
import os
import sys
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gen_luna_data as G

def write_optional_manifest():
    """기존 로더와 무관한 선택적 배포 목록. 파일 누락·QA 혼입을 검수할 때 사용한다."""
    jdir = os.path.join(G.OUT, "json")
    names = []
    for root, _, files in os.walk(jdir):
        for filename in files:
            path = os.path.join(root, filename)
            rel = os.path.relpath(path, jdir).replace(os.sep, "/")
            if filename.endswith(".json") and rel != "manifest.json":
                names.append(rel)
    names.sort()
    metadata = {}
    for name in names:
        with open(os.path.join(jdir, *name.split("/")), "rb") as f:
            payload = f.read()
        metadata[name] = {"sha256": hashlib.sha256(payload).hexdigest(), "bytes": len(payload)}
    material = "".join(f"{name}\0{metadata[name]['sha256']}\0" for name in names).encode("utf-8")
    qa_files = [name for name in names if name == "script/bar/day99.json" or name.startswith("script/qa/") or name.startswith("qa/")]
    manifest = {
        "bundle_contract_version": "1.0.0",
        "data_schema_version": dict((k, v) for k, v, _ in G.CONFIG)["data_schema_version"],
        "bundle_sha256": hashlib.sha256(material).hexdigest(),
        "production_files": [name for name in names if name not in qa_files],
        "qa_files": qa_files,
        "compatibility": {
            "preferred_cocktail_fields": ["target_mix_method", "target_prep_action", "tags[].id"],
            "legacy_cocktail_fields": ["mix", "prep", "tags[].ko"],
        },
        "files": metadata,
    }
    with open(os.path.join(jdir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

# Config 값 타입은 시트의 type 컬럼이 정본 (v2.2) — 엑셀이 3.0을 3으로 뭉개도 되돌린다.
# 시드 키 기억 방식(SEED_FLOAT_CONFIG)은 신규 키를 못 지켜서 폐지.
def _config_cast(v, t):
    if t == "float" and isinstance(v, int) and not isinstance(v, bool): return float(v)
    if t == "int" and isinstance(v, float) and v == int(v): return int(v)
    if t == "bool" and isinstance(v, str): return v.upper() == "TRUE"
    if t == "str": return "" if v is None else str(v)
    return v

# 컬럼 타입 명세 — 엑셀 왕복 시 타입이 뭉개지는 컬럼만 명시 (그 외는 그대로)
NULLABLE = {   # 빈 칸을 ""가 아니라 None(널)로
    "Cocktails": {"glass", "garnish", "color2"},
    "RecipeLines": {"qty", "unit"},          # fill_up·수량 tbd 라인은 수량이 널
    "ScoreBands": {"max_ratio"},             # 마지막 구간은 상한 없음(널)
    "Scenes": {"day", "when", "group", "start_mode", "on_complete_effects"},
    "ShelfItems": {"shop_price", "category", "color", "sprite", "prep_action", "default_target_qty", "default_target_unit", "shelf_group", "liquid_alpha"},
    "Characters": {"alive_flag", "enter_sfx", "exit_sfx"},
    "RandomWaves": {"order"},
    "InteractPoints": {"facing", "spawn_when", "interact_when", "action_type", "action_ref"},
    "Transitions": {"target_spot"},
    "OrderRules": {"when", "effects", "react"},
    "Endings": {"when"},
}
BOOL_COLS = {  # TRUE/FALSE 문자열 → 불리언
    "Scenes": {"skippable"},
    "RegularSlots": {"branch_choice", "must_serve"},
    "RandomWaves": {"branch_choice"},
    "RecipeLines": {"is_core", "auto_apply", "scored"},
    "ScoreBands": {"min_inclusive", "max_inclusive"},
    "Characters": {"affinity"},
    "GuestBodies": {"is_default"},
}
FLOAT_COLS = { # 8.0이 8로 읽히는 문제 → float 강제
    "Cocktails": {"abv"},
    "SettlementRules": {"sale_rate", "tip_rate", "refund_rate"},
    "ScoreBands": {"min_ratio", "max_ratio"},
    "Personalities": {"tip_mult", "patience_mult", "think_chance"},
    "ShelfItems": {"default_target_qty", "liquid_alpha"},
}

# 시트 → (전역 이름, 읽을 컬럼 수). Cocktails의 (파생) 컬럼과 Characters의 base_body는 별도 처리.
SHEET_SPEC = {
    "Cocktails": ("COCKTAILS", len(G.CK_COLS)),
    "RecipeLines": (None, 9),
    "ShelfItems": ("SHELF_ITEMS", len(G.SHELF_COLS)),
    "Characters": (None, len(G.CHAR_COLS) + 1),
    "Expressions": ("EXPRESSIONS", len(G.EXPR_COLS)),
    "ExpressionParts": ("EXPRESSION_PARTS", len(G.EXPRPART_COLS)),
    "Cutscenes": ("CUTSCENES", len(G.CUT_COLS)),
    "ResourceMap": ("RESOURCE_MAP", len(G.RESMAP_COLS)),
    "FieldAnims": ("FIELD_ANIMS", len(G.FIELD_COLS)),
    "Personalities": ("PERSONALITIES", len(G.PERS_COLS)),
    # 기존 4열 태그 계약은 유지하고, 5번째 tag_id만 빌드에서 별도 흡수한다.
    "Tags": (None, len(G.TAG_COLS) + 1),
    "GuestBodies": ("GUEST_BODIES", len(G.GBODY_COLS)),
    "GuestBodyExclusions": ("GUEST_BODY_EXCLUSIONS", len(G.GBEXCL_COLS)),
    "Barks": ("BARKS", len(G.BARK_COLS)),
    "Tastes": ("TASTES", len(G.TASTE_COLS)),
    "Dossier": ("DOSSIER", len(G.DOSSIER_COLS)),
    "Days": ("DAYS", len(G.DAY_COLS)),
    "RandomWaves": ("RANDOM_WAVES", len(G.WAVE_COLS)),
    "RegularSlots": ("REGULAR_SLOTS", len(G.RSLOT_COLS)),
    "Spots": ("SPOTS", len(G.SPOT_COLS)),
    "InteractPoints": ("POINTS", len(G.POINT_COLS)),
    "Transitions": ("TRANSITIONS", len(G.TRANSITION_COLS)),
    "Scenes": ("SCENES", len(G.SCENE_COLS)),
    "Steps": ("STEPS", len(G.STEP_COLS)),
    "Choices": ("CHOICES", len(G.CHOICE_COLS)),
    "OrderRules": ("ORDERS", len(G.ORDER_COLS)),
    "Quests": ("QUESTS", len(G.QUEST_COLS)),
    "QuestStages": ("QUEST_STAGES", len(G.QSTAGE_COLS)),
    "Endings": ("ENDINGS", len(G.END_COLS)),
    "Config": ("CONFIG", len(G.CONFIG_COLS)),
    "GradeCuts": ("GRADE_CUTS", 2),
    "SettlementRules": ("SETTLEMENT_RULES", len(G.SETTLE_COLS)),
    "ScoreBands": ("SCORE_BANDS", len(G.SCOREBAND_COLS)),
    "AffinityMatrix": ("AFFINITY_MATRIX", 6),
    "UIStrings": ("UI_STRINGS", len(G.UI_COLS)),
    "TextTags": ("TEXT_TAGS", len(G.TEXTTAG_COLS)),
    "BarkSituations": ("BARK_SITUATIONS", len(G.BARKSIT_COLS)),
}
REQUIRED = set(SHEET_SPEC)


def rows_of(ws, ncols, sheet):
    """시트 → 튜플 리스트. 완전 빈 행은 스킵, 빈 셀은 ""(NULLABLE 컬럼만 None)."""
    nul = NULLABLE.get(sheet, set())
    bools = BOOL_COLS.get(sheet, set())
    floats = FLOAT_COLS.get(sheet, set())
    hdr = [c.value for c in ws[1][:ncols]]
    hdr += [None] * (ncols - len(hdr))
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        vals = r[:ncols]
        if all(v is None or v == "" for v in vals):
            continue
        row = []
        for i in range(ncols):
            v = vals[i] if i < len(vals) else None
            col = hdr[i]
            if v is None:
                v = None if col in nul else ""
            elif isinstance(v, str):
                v = v.strip()
                # Excel/편집 도구에 따라 빈 셀이 None 또는 ""로 왕복된다.
                # nullable 컬럼은 두 표현을 같은 null로 읽어 데이터 의미를 보존한다.
                if not v and col in nul:
                    v = None
                elif col in bools:
                    v = v.upper() == "TRUE"
            if col in bools and isinstance(v, (int, float)) and not isinstance(v, bool):
                v = bool(v)
            if col in floats and isinstance(v, int) and not isinstance(v, bool):
                v = float(v)
            row.append(v)
        out.append(tuple(row))
    return out


def collect(paths):
    """파일들을 읽어 시트를 수집. v2.0: 한 시트는 한 파일에만 — 중복이면 dup에 기록(빌드 에러).
    참고용으로 남의 시트를 복사해 두면 행이 조용히 섞여 '유령 데이터'가 되므로 병합하지 않는다."""
    tables, origin, unknown, dup = {}, {}, [], []
    for p in paths:
        wb = load_workbook(p, data_only=True)
        base = os.path.basename(p)
        for name in wb.sheetnames:
            if name == "INFO":
                continue
            if name not in SHEET_SPEC:
                unknown.append(f"{base}:{name}")
                continue
            if name in origin:
                dup.append(f"{name} ({origin[name]} ↔ {base})")
                continue
            origin[name] = base
            tables[name] = rows_of(wb[name], SHEET_SPEC[name][1], name)
    return tables, origin, unknown, dup


def load_into_globals(tables):
    # 특수 1 — RecipeLines → RECIPES 딕셔너리 (레시피 없는 칵테일도 빈 키 유지)
    # 시트의 is_core/auto_apply는 시드의 플래그 문자열로 복원. scored는 파생 참고 컬럼 —
    # 플래그로 설명되지 않는 scored=FALSE만 noscore로 되살린다(현재 실사용 0 — 예약).
    G.COCKTAILS = tables["Cocktails"]
    rec = {}
    for cid, seq, action, ing, qty, unit, is_core, auto_apply, scored in sorted(
            tables["RecipeLines"], key=lambda x: (str(x[0]), x[1])):
        flags = [f for f, on in (("core", is_core), ("auto", auto_apply)) if on]
        if not scored and not auto_apply:
            flags.append("noscore")
        rec.setdefault(cid, []).append((action, ing, qty, unit, "|".join(flags)))
    for c in G.COCKTAILS:
        rec.setdefault(c[0], [])
    G.RECIPES = rec
    # 특수 2 — Characters: 마지막 base_body 컬럼 → BASE_BODY
    ch = tables["Characters"]
    G.CHARACTERS = [tuple(r[:len(G.CHAR_COLS)]) for r in ch]
    G.BASE_BODY = {r[0]: r[len(G.CHAR_COLS)] for r in ch if r[len(G.CHAR_COLS)]}
    # 특수 3 — Tags: 기존 한글 정본 4열과 호환하면서 불변 tag_id 5열을 병행한다.
    tag_rows = tables["Tags"]
    G.TAGS = [tuple(r[:len(G.TAG_COLS)]) for r in tag_rows]
    G.TAG_EN = {r[0]: r[1] for r in tag_rows}
    G.TAG_CATEGORY = {r[0]: r[2] for r in tag_rows}
    G.TAG_IDS = {r[0]: r[len(G.TAG_COLS)] for r in tag_rows}
    # 나머지 1:1
    for name, (gname, _) in SHEET_SPEC.items():
        if gname:
            setattr(G, gname, tables[name])
    # 특수 4 — Config 값 타입 복원: 시트의 type 컬럼대로 (4컬럼 → 내부 3튜플)
    G.CONFIG = [(k, _config_cast(v, t), n) for k, v, t, n in G.CONFIG]


def apply_minimal_street_scope():
    """데모 거리 판정을 최근접 1개로 단순화하고 priority 전용 QA만 제외한다.

    엑셀은 수정하지 않는다는 작업 규칙 때문에 제거 대상으로 확정된 QA 행을
    빌드 입력에서 제외한다. priority 필드는 호환용으로 남지만 런타임 판정에는
    사용하지 않는다.
    """
    excluded_points = {"p_qa_priority_low", "p_qa_priority_high"}
    excluded_scenes = {"qa_priority_low", "qa_priority_high"}
    G.POINTS = [row for row in G.POINTS if row[0] not in excluded_points]
    G.SCENES = [row for row in G.SCENES if row[0] not in excluded_scenes]
    G.STEPS = [row for row in G.STEPS if row[0] not in excluded_scenes]


def main():
    args = sys.argv[1:]
    if "--strict" in args:
        G.STRICT_L10N = True   # 영어 누락을 에러로 승격 (8월 번역 완료 후)
        args.remove("--strict")
    paths = args if args else [os.path.join(G.OUT, f) for f in ("LUNA_System.xlsx", "LUNA_Narrative.xlsx")]
    src = " + ".join(os.path.basename(p) for p in paths)
    for p in paths:
        if not os.path.exists(p):
            print(f"❌ 시트 파일 없음: {p}"); sys.exit(1)

    tables, origin, unknown, dup = collect(paths)
    if dup:
        print("❌ 같은 시트가 두 파일에 존재 — 한 시트는 한 파일에만 있어야 합니다 (참고용 복사 금지):")
        for d in dup:
            print(f"   · {d}")
        sys.exit(1)
    missing = REQUIRED - set(origin)
    if missing:
        print(f"❌ 시트 누락: {', '.join(sorted(missing))}"); sys.exit(1)

    load_into_globals(tables)
    apply_minimal_street_scope()
    derived = G.derive()
    errors, report = G.validate(derived)
    errors += G.validate_street_runtime_contract(G.build_street_runtime_contract())

    lines = [f"===== L.U.N.A 시트 빌드 리포트 (build.py · 원본: {src}) ====="]
    lines.append("· 거리 v2.6.0 대화 계약 — InteractPoints가 dialogue_flows를 소유하고 street는 대본 스텝만 배포")
    lines.append("· 거리 대상 선정 — 가장 가까운 유효 대상 우선, 완전히 동률이면 point.id 오름차순·priority 전용 Day 99 QA 2종 제외")
    if unknown:
        lines.append(f"⚠ 알 수 없는 시트(무시됨): {', '.join(unknown)}")
    by_file = {}   # 파일별 시트·행수 — "남의 파일 안 당겨오고 빌드" 사고를 눈으로 잡는 용도
    for name, base in origin.items():
        by_file.setdefault(base, []).append(f"{name}({len(tables[name])})")
    for base, items in sorted(by_file.items()):
        lines.append(f"· {base} — 시트 {len(items)}개: " + ", ".join(sorted(items)))
    lines += report + [""]
    if errors:
        lines.append(f"❌ 오류 {len(errors)}건 — JSON 미출력:")
        lines += ["  " + e for e in errors]
    else:
        lines.append("✅ 검증 통과 — 참조 무결성 · 파생 규칙 · DSL 문법 · 플래그 · L10N · 고정 대사 ID · 루나 대사 규칙")
    lines.append(f"칵테일 {len(G.COCKTAILS)} / 선반 {len(G.SHELF_ITEMS)} / 캐릭터 {len(G.CHARACTERS)} / "
                 f"씬 {len(G.SCENES)} / 스텝 {len(G.STEPS)} / 선택지 {len(G.CHOICES)} / "
                 f"웨이브 {len(G.RANDOM_WAVES)} / 단골슬롯 {len(G.REGULAR_SLOTS)}")
    print("\n".join(lines))
    with open(os.path.join(G.OUT, "빌드리포트.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    if errors:
        sys.exit(1)
    G.emit_json(derived)
    write_optional_manifest()
    print("→ json/ 출력 완료 (시트가 원본입니다 — xlsx는 재생성하지 않음)")


if __name__ == "__main__":
    main()
